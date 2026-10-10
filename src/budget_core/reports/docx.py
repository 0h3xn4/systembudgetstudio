"""DOCX renderer (python-docx). Identical documents give identical bytes: the core properties
carry the provenance time, and the ZIP is rewritten with fixed timestamps and order."""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from budget_core.plots.render import image_size
from budget_core.reports.document import NA, Column, Figure, ReportDocument, Section, Table
from budget_core.reports.document import format_cell as _format_cell
from budget_core.reports.textsafe import clean_document
from budget_core.reports.zipnorm import normalise_zip

FONT = "IBM Plex Sans"
HEADER_FILL = "E0E0E0"
EMPHASIS_FILL = "F4F4F4"
RED = RGBColor(0xDA, 0x1E, 0x28)
PAGE_W_MM, PAGE_H_MM, MARGIN_MM = 297.0, 210.0, 12.0


def _shade(cell: object, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def _run(paragraph: Paragraph, text: str, *, bold: bool = False, size: float = 8.0) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)


CELL_STYLES = (
    # (name, bold, right aligned)
    ("Budget cell", False, False),
    ("Budget cell bold", True, False),
    ("Budget cell right", False, True),
    ("Budget cell right bold", True, True),
)


def _cell_styles(doc: object) -> dict[tuple[bool, bool], str]:
    """Paragraph styles for table cells, so a cell needs one run with text only (setting the
    font on every run of a large table is the slowest part of the DOCX export)."""
    styles = doc.styles  # type: ignore[attr-defined]
    found: dict[tuple[bool, bool], str] = {}
    for name, bold, right in CELL_STYLES:
        try:
            style = styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles["Normal"]
            style.font.name = FONT
            style.font.size = Pt(7.5)
            style.font.bold = bold
            if right:
                style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        found[(bold, right)] = style.style_id
    return found


def _table(doc: object, table: Table) -> None:
    heading = doc.add_paragraph()  # type: ignore[attr-defined]
    _run(heading, table.title, bold=True, size=9.0)
    heading.paragraph_format.keep_with_next = True
    grid = doc.add_table(rows=1, cols=len(table.columns))  # type: ignore[attr-defined]
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    styles = _cell_styles(doc)
    for cell, column in zip(grid.rows[0].cells, table.columns, strict=True):
        paragraph = cell.paragraphs[0]
        paragraph._p.style = styles[(True, False)]  # the style id; no lookup by name
        paragraph.add_run(column.header)
        _shade(cell, HEADER_FILL)
    last = len(table.rows) - 1
    for i, values in enumerate(table.rows):
        cells = grid.add_row().cells
        emphasise = table.emphasise_last_row and i == last
        for cell, value, column in zip(cells, values, table.columns, strict=True):
            paragraph = cell.paragraphs[0]
            paragraph._p.style = styles[(emphasise, column.kind != "text")]
            paragraph.add_run(_format(value, column))
            if emphasise:
                _shade(cell, EMPHASIS_FILL)
    # repeat the header row on every page
    header_props = grid.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    if table.note:
        note = doc.add_paragraph()  # type: ignore[attr-defined]
        _run(note, table.note, size=7.5)
    doc.add_paragraph()  # type: ignore[attr-defined]


def _format(value: object, column: Column) -> str:
    if value is None:
        return NA
    return _format_cell(value, column)  # type: ignore[arg-type]


def _figure(doc: object, figure: Figure) -> None:
    heading = doc.add_paragraph()  # type: ignore[attr-defined]
    _run(heading, figure.title, bold=True, size=9.0)
    heading.paragraph_format.keep_with_next = True
    width_px, height_px = image_size(figure.png)
    max_w, max_h = PAGE_W_MM - 2 * MARGIN_MM, PAGE_H_MM - 2 * MARGIN_MM - 30.0
    scale = min(max_w / width_px, max_h / height_px)
    doc.add_picture(io.BytesIO(figure.png), width=Mm(width_px * scale))  # type: ignore[attr-defined]
    if figure.alt:
        caption = doc.add_paragraph()  # type: ignore[attr-defined]
        _run(caption, figure.alt, size=7.5)


def _section(doc: object, section: Section) -> None:
    heading = doc.add_paragraph()  # type: ignore[attr-defined]
    _run(heading, section.title, bold=True, size=12.0)
    heading.paragraph_format.keep_with_next = True
    heading.style = doc.styles["Heading 1"]  # type: ignore[attr-defined]
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x16, 0x16, 0x16)
    for text in section.paragraphs:
        _run(doc.add_paragraph(), text)  # type: ignore[attr-defined]
    for table in section.tables:
        _table(doc, table)
    for figure in section.figures:
        _figure(doc, figure)


def render_docx(doc: ReportDocument) -> bytes:
    doc = clean_document(doc)
    document = Document()
    page = document.sections[0]
    page.orientation = WD_ORIENT.LANDSCAPE
    page.page_width, page.page_height = Mm(PAGE_W_MM), Mm(PAGE_H_MM)
    page.left_margin = page.right_margin = Mm(MARGIN_MM)
    page.top_margin = page.bottom_margin = Mm(MARGIN_MM)
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)

    title = document.add_paragraph()
    _run(title, doc.title, bold=True, size=16.0)
    if doc.banner:
        banner = document.add_paragraph()
        _run(banner, doc.banner, bold=True, size=9.0)
        for run in banner.runs:
            run.font.color.rgb = RED
    for section in doc.sections:
        _section(document, section)

    prov = doc.provenance
    footer = page.footer.paragraphs[0]
    _run(
        footer,
        f"{prov.tool} {prov.tool_version} | {prov.project_name} rev {prov.project_revision} | "
        f"{prov.scenario} | {prov.generated} | {prov.user}",
        size=7.0,
    )

    stamp = datetime.strptime(prov.generated, "%Y-%m-%dT%H:%M:%SZ")
    props = document.core_properties
    props.author = prov.user
    props.last_modified_by = prov.tool
    props.title = doc.title
    props.subject = prov.project_name
    props.created = stamp
    props.modified = stamp
    props.revision = 1
    props.comments = ""
    props.keywords = ""
    props.category = ""
    buffer = io.BytesIO()
    document.save(buffer)
    return normalise_zip(buffer.getvalue())
