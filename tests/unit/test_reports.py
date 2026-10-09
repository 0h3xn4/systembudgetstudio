"""Report content, provenance and byte-identical regeneration (spec constraints 14 and 19)."""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from budget_core.io.project_loader import load_project
from budget_core.power.static_budget import static_power_budget
from budget_core.problems import Problem
from budget_core.provenance import Provenance, make_provenance
from budget_core.reports.document import NA, Column, ReportDocument, format_cell
from budget_core.reports.export import result_csv, result_json
from budget_core.reports.pdf import render_pdf
from budget_core.reports.power_report import build_power_report
from budget_core.reports.xlsx import render_xlsx
from budget_core.reports.zipnorm import normalise_zip
from tests.helpers import write_valid_project

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
WHEN = datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)


def prov(project_root: Path = EXAMPLES / "cubesat_3u") -> Provenance:
    project = load_project(project_root).project
    assert project is not None
    return make_provenance(project, user="Test User", generated_at=WHEN)


def document(name: str = "cubesat_3u") -> ReportDocument:
    result = load_project(EXAMPLES / name)
    assert result.project is not None
    power = static_power_budget(result.project)
    return build_power_report(result.project, power, prov(EXAMPLES / name), result.problems)


def test_provenance_fields() -> None:
    p = prov()
    assert p.tool == "System Budget Studio" and p.user == "Test User"
    assert p.generated == "2026-01-02T03:04:05Z"
    assert p.project_revision == "1" and p.scenario.startswith("static")
    assert {n for n, _ in p.libraries} >= {"pydantic", "openpyxl", "reportlab"}


def test_provenance_uses_source_date_epoch_and_budget_user(monkeypatch: pytest.MonkeyPatch) -> None:
    project = load_project(EXAMPLES / "cubesat_3u").project
    assert project is not None
    monkeypatch.setenv("SOURCE_DATE_EPOCH", "86400")
    monkeypatch.setenv("BUDGET_USER", "Env User")
    p = make_provenance(project)
    assert p.generated == "1970-01-02T00:00:00Z" and p.user == "Env User"


def test_format_cell() -> None:
    assert format_cell(None, Column("x", "number")) == NA
    assert format_cell(1.23456, Column("x", "number", 2)) == "1.23"
    assert format_cell(0.1, Column("x", "percent", 3)) == "10.0 %"
    assert format_cell("abc", Column("x")) == "abc"


def test_document_structure_and_unique_sheet_names() -> None:
    doc = document()
    names = [s.sheet_name for s in doc.sections]
    assert names[0] == "Summary" and names[-3:] == ["Assumptions", "Problems", "Provenance"]
    assert len(set(names)) == len(names) and all(len(n) <= 31 for n in names)
    assert len(doc.sections) == 1 + 5 + 3  # summary, 5 spacecraft modes, three closing sections


def test_incomplete_banner_when_inputs_are_placeholders() -> None:
    doc = document()
    assert doc.banner.startswith("INCOMPLETE")
    summary = doc.sections[0].tables[0]
    assert all(cell in (NA, None) for cell in summary.rows[0][2:6])  # margined and source columns


def test_no_banner_and_numbers_when_everything_is_sourced(tmp_path: Path) -> None:
    write_valid_project(tmp_path)
    result = load_project(tmp_path)
    assert result.project is not None
    power = static_power_budget(result.project)
    doc = build_power_report(result.project, power, prov(tmp_path), result.problems)
    assert doc.banner == ""
    nominal = next(s for s in doc.sections if s.sheet_name == "Mode nominal")
    totals = nominal.tables[-1]
    assert all(isinstance(r[1], float) for r in totals.rows)


def test_xlsx_content_and_properties() -> None:
    doc = document()
    wb = load_workbook(io.BytesIO(render_xlsx(doc)))
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Provenance"
    assert wb.properties.creator == "System Budget Studio"
    assert wb.properties.created == datetime(2026, 1, 2, 3, 4, 5)
    text = "\n".join(
        str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert "INCOMPLETE" in text and "n/a" in text
    assert "Test User" in text and "2026-01-02T03:04:05Z" in text and "PWR-SOURCE" in text
    assert "SOURCE_MISSING" in text  # equations without a source are flagged


def test_xlsx_numbers_are_numbers_with_formats(tmp_path: Path) -> None:
    write_valid_project(tmp_path)
    result = load_project(tmp_path)
    assert result.project is not None
    doc = build_power_report(
        result.project, static_power_budget(result.project), prov(tmp_path), result.problems
    )
    ws = load_workbook(io.BytesIO(render_xlsx(doc)))["Mode nominal"]
    numeric = [c for row in ws.iter_rows() for c in row if isinstance(c.value, float)]
    assert numeric and all(c.number_format.startswith("0") for c in numeric)


def test_outputs_are_byte_identical_in_process() -> None:
    doc = document()
    assert render_xlsx(doc) == render_xlsx(doc)
    assert render_pdf(doc) == render_pdf(doc)


def test_pdf_basics() -> None:
    data = render_pdf(document())
    assert data.startswith(b"%PDF-") and b"IBMPlexSans" in data and len(data) > 5000


SCRIPT = """
import sys
from datetime import UTC, datetime
from pathlib import Path
from budget_core.io.project_loader import load_project
from budget_core.power.static_budget import static_power_budget
from budget_core.provenance import make_provenance
from budget_core.reports.power_report import build_power_report
from budget_core.reports.pdf import render_pdf
from budget_core.reports.xlsx import render_xlsx
root = Path(sys.argv[1])
r = load_project(root)
p = make_provenance(r.project, user="U", generated_at=datetime(2026, 1, 2, tzinfo=UTC))
doc = build_power_report(r.project, static_power_budget(r.project), p, r.problems)
Path(sys.argv[2]).write_bytes(render_xlsx(doc))
Path(sys.argv[3]).write_bytes(render_pdf(doc))
"""


def test_regeneration_is_byte_identical_across_processes(tmp_path: Path) -> None:
    outs = []
    for i, seed in enumerate(("1", "12345")):
        xlsx, pdf = tmp_path / f"a{i}.xlsx", tmp_path / f"a{i}.pdf"
        env = {**os.environ, "PYTHONHASHSEED": seed, "QT_QPA_PLATFORM": "offscreen"}
        subprocess.run(
            [sys.executable, "-c", SCRIPT, str(EXAMPLES / "microsat_150kg"), str(xlsx), str(pdf)],
            check=True,
            env=env,
        )
        outs.append((xlsx.read_bytes(), pdf.read_bytes()))
    assert outs[0] == outs[1]


def test_normalise_zip_is_stable_and_valid() -> None:
    raw = render_xlsx(document())
    assert normalise_zip(raw) == normalise_zip(raw)


def test_json_export_is_deterministic_and_complete() -> None:
    result = load_project(EXAMPLES / "cubesat_3u")
    assert result.project is not None
    power = static_power_budget(result.project)
    text = result_json(power, prov())
    assert text == result_json(power, prov())
    data = json.loads(text)
    assert data["provenance"]["user"] == "Test User"
    assert [m["mode_id"] for m in data["modes"]] == sorted(m["mode_id"] for m in data["modes"])
    assert data["modes"][0]["average"]["system_margined_w"] is None  # placeholders -> null
    assert data["assumptions"] and text.endswith("\n")


def test_csv_export_has_header_provenance_comment_free_rows() -> None:
    result = load_project(EXAMPLES / "cubesat_3u")
    assert result.project is not None
    power = static_power_budget(result.project)
    text = result_csv(power.modes[0])
    lines = text.strip().split("\n")
    assert lines[0].startswith("unit_id,unit_name,subsystem,bus,power_mode,avg_power_w")
    assert len(lines) == 1 + 6 and "\r" not in text and ",," in text  # empty cell for n/a


def test_problems_listed_in_report() -> None:
    doc = document()
    section = next(s for s in doc.sections if s.sheet_name == "Problems")
    codes = {str(r[1]) for r in section.tables[0].rows}
    assert "CONFIG_PLACEHOLDER" in codes and "RESULT_INCOMPLETE" in codes
    assert isinstance(Problem, type)


def test_xlsx_does_not_depend_on_the_clock() -> None:
    """openpyxl stamps the save time into docProps/core.xml; renders must not carry it."""
    import time

    doc = document()
    first = render_xlsx(doc)
    time.sleep(1.2)
    assert render_xlsx(doc) == first
    wb = load_workbook(io.BytesIO(first))
    assert wb.properties.modified == datetime(2026, 1, 2, 3, 4, 5)


def test_pdf_does_not_depend_on_the_clock() -> None:
    import time

    doc = document()
    first = render_pdf(doc)
    time.sleep(1.2)
    assert render_pdf(doc) == first
