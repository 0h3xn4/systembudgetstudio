"""Output safety: file names built from project text stay inside the output folder, writes do not
follow symlinks, and text that a spreadsheet would run as a formula or that XML cannot hold is
neutralised."""

from __future__ import annotations

import csv
import io
import os
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from budget_core.provenance import Provenance
from budget_core.reports.csvutil import csv_writer
from budget_core.reports.document import Column, ReportDocument, Section, Table
from budget_core.reports.files import OutputError, safe_name, write_file
from budget_core.reports.run import BudgetOutput, write_outputs

PROV = Provenance("t", "0", "3", (), "P", "1", "s", "2026-01-01T00:00:00Z", "u")


# ---- file names -----------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("mass_static_launch.csv", "mass_static_launch.csv"),
        ("mass_static_a/b.csv", "mass_static_a_b.csv"),
        ("../../x.csv", "_.._x.csv"),
        ("a\\b:c*.csv", "a_b_c_.csv"),
        (".hidden", "hidden"),
        ("con.csv", "_con.csv"),
        ("NUL", "_NUL"),
        ("", "_"),
        ("tab\tname.csv", "tab_name.csv"),
        ("é-ok.csv", "é-ok.csv"),
    ],
)
def test_safe_name(raw: str, expected: str) -> None:
    assert safe_name(raw) == expected


def test_long_names_are_shortened_keeping_the_extension() -> None:
    name = safe_name("x" * 400 + ".csv")
    assert len(name) <= 120 and name.endswith(".csv")


def test_write_file_stays_inside_the_folder(tmp_path: Path) -> None:
    out = tmp_path / "out"
    path = write_file(out, "../../escape.csv", "x")
    assert path.parent == out and "/" not in path.name and "\\" not in path.name
    assert not (tmp_path / "escape.csv").exists()


@pytest.mark.skipif(sys.platform == "win32", reason="symlinks need privileges on Windows")
def test_write_file_replaces_a_symlink_instead_of_following_it(tmp_path: Path) -> None:
    victim = tmp_path / "victim.txt"
    victim.write_text("precious", encoding="utf-8")
    out = tmp_path / "out"
    out.mkdir()
    os.symlink(victim, out / "a.csv")
    write_file(out, "a.csv", "new")
    assert victim.read_text(encoding="utf-8") == "precious"
    assert (out / "a.csv").read_text(encoding="utf-8") == "new"
    assert not (out / "a.csv").is_symlink()


def test_write_file_leaves_no_partial_file_on_failure(tmp_path: Path) -> None:
    out = tmp_path / "out"
    with pytest.raises(TypeError):
        write_file(out, "a.csv", 123)  # type: ignore[arg-type]
    assert not list(out.glob("*")) if out.exists() else True


def test_colliding_names_are_refused_not_overwritten(tmp_path: Path) -> None:
    doc = ReportDocument("R", PROV, ())
    one = BudgetOutput("p", doc, None, PROV, (("a_b.csv", "1"), ("a/b.csv", "2")), "{}")
    with pytest.raises(OutputError) as info:
        write_outputs([one], tmp_path / "o", {"csv"})
    assert "same file name" in str(info.value)
    assert not (tmp_path / "o" / "a_b.csv").exists()


def test_csv_names_from_project_text_cannot_escape(tmp_path: Path) -> None:
    doc = ReportDocument("R", PROV, ())
    out = BudgetOutput("p", doc, None, PROV, (("mass_static_x/../../../PWN.csv", "x"),), "{}")
    written = write_outputs([out], tmp_path / "o", {"csv"})
    assert all(p.parent == tmp_path / "o" for p in written)
    assert not (tmp_path / "PWN.csv").exists()


# ---- CSV text --------------------------------------------------------------------------------
def test_csv_text_that_a_spreadsheet_would_run_is_defused() -> None:
    buffer = io.StringIO()
    writer = csv_writer(buffer)
    writer.writerow(
        ["=HYPERLINK(1)", "+1+1", "@SUM(A1)", "-cmd|x", "\tx", "ok", "-0.5", "+3", "1e-9"]
    )
    row = next(csv.reader(io.StringIO(buffer.getvalue())))
    assert row[:5] == ["'=HYPERLINK(1)", "'+1+1", "'@SUM(A1)", "'-cmd|x", "'\tx"]
    assert row[5:] == ["ok", "-0.5", "+3", "1e-9"]  # numbers and plain text are untouched


@pytest.mark.parametrize("text", ["+2.500 W", "-0.400 W", "+20.0 %", "-10.0 pp", "+0.5 dB"])
def test_signed_numbers_with_a_unit_are_not_defused(text: str) -> None:
    from budget_core.reports.csvutil import defuse

    assert defuse(text) == text


@pytest.mark.parametrize("text", ["+1+1", "-1+cmd|' /C calc'!A0", "+1 cmd|x", "=1+1", "@A1"])
def test_formula_like_text_is_defused(text: str) -> None:
    from budget_core.reports.csvutil import defuse

    assert defuse(text) == "'" + text


def test_csv_writer_passes_non_strings_through() -> None:
    buffer = io.StringIO()
    csv_writer(buffer).writerow([1, 2.5, None])
    assert buffer.getvalue() == "1,2.5,\n"


# ---- XLSX and DOCX text ---------------------------------------------------------------------
def doc_with(text: str) -> ReportDocument:
    table = Table("T", (Column("Name"),), ((text,),))
    return ReportDocument(text, PROV, (Section(text, "S", paragraphs=(text,), tables=(table,)),))


def test_xlsx_has_no_formula_cells_from_project_text() -> None:
    from budget_core.reports.xlsx import render_xlsx

    wb = load_workbook(io.BytesIO(render_xlsx(doc_with('=HYPERLINK("http://x","c")'))))
    kinds = {c.data_type for ws in wb for row in ws.iter_rows() for c in row if c.value}
    assert "f" not in kinds


@pytest.mark.parametrize("kind", ["xlsx", "docx", "pdf"])
def test_control_characters_do_not_crash_renderers(kind: str) -> None:
    from budget_core.reports.docx import render_docx
    from budget_core.reports.pdf import render_pdf
    from budget_core.reports.xlsx import render_xlsx

    render = {"xlsx": render_xlsx, "docx": render_docx, "pdf": render_pdf}[kind]
    assert render(doc_with("A\x07B\x0bC\x00D")).startswith((b"PK", b"%PDF"))
