"""Time-domain power outputs: report document, plots, XLSX/PDF/DOCX, CSV, JSON, CLI."""

from __future__ import annotations

import io
import json
import shutil
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from budget_cli.main import main
from budget_core.plots.power import power_plot
from budget_core.plots.render import image_size, render_png
from budget_core.power.time_domain import time_domain_budget
from budget_core.provenance import make_provenance
from budget_core.reports.docx import render_docx
from budget_core.reports.pdf import render_pdf
from budget_core.reports.run import timeline_output, write_outputs
from budget_core.reports.timeline_export import series_csv
from budget_core.reports.xlsx import render_xlsx
from tests.power_helpers import synthetic_env, synthetic_run, td_project

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
WHEN = datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)
FIXED = ["--user", "Test User", "--date", "2026-01-02T03:04:05Z"]


@pytest.fixture(scope="module")
def output():  # type: ignore[no-untyped-def]
    project = td_project()
    env = synthetic_env(12000.0, 20.0, [(3000.0, 4800.0), (9000.0, 10800.0)])
    result = time_domain_budget(project, synthetic_run(env, default_mode="b"))
    provenance = make_provenance(project, scenario="synthetic", user="Test User", generated_at=WHEN)
    return result, timeline_output(project, result, provenance)


def test_plot_png_is_a_png_with_one_panel_per_series_group(output) -> None:  # type: ignore[no-untyped-def]
    result, _ = output
    spec = power_plot(result, "bol")
    assert [p.title for p in spec.panels] == ["Power", "Battery", "Margin"]
    assert [b.label for b in spec.bands] == ["Eclipse"]
    png = render_png(spec)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    width, height = image_size(png)
    assert width == 1400 and height > 800
    assert render_png(spec) == png  # deterministic


def test_plot_without_battery_inputs_has_no_battery_panel() -> None:
    project = td_project()
    system = project.config.power_system
    assert system is not None
    env = synthetic_env(1200.0, 20.0)
    gap = system.battery.model_copy(
        update={
            "charge_efficiency_ratio": system.battery.charge_efficiency_ratio.model_copy(
                update={"value": None}
            )
        }
    )
    project = type(project)(
        **{
            **project.__dict__,
            "config": type(project.config)(
                margin_policy=project.config.margin_policy,
                power_config=project.config.power_config,
                power_system=system.model_copy(update={"battery": gap}),
            ),
        }
    )
    result = time_domain_budget(project, synthetic_run(env))
    assert [p.title for p in power_plot(result, "bol").panels] == ["Power", "Margin"]


def test_report_sections_and_content(output) -> None:  # type: ignore[no-untyped-def]
    _, out = output
    titles = [s.title for s in out.document.sections]
    assert titles == [
        "Summary",
        "Plots, BOL",
        "Plots, EOL",
        "Modes and orbits",
        "Static power",
        "Assumptions",
        "Problems",
        "Provenance",
    ]
    summary = out.document.sections[0]
    assert any("Eclipse fraction" in p for p in summary.paragraphs)
    violations = next(t for t in summary.tables if t.title == "Violations")
    findings = {row[1] for row in violations.rows}
    assert "Negative energy balance over an orbit" in findings
    assert out.document.banner == ""  # every input is given


def test_xlsx_has_the_figure_and_the_tables(output) -> None:  # type: ignore[no-untyped-def]
    _, out = output
    data = render_xlsx(out.document)
    assert render_xlsx(out.document) == data
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert any(n.startswith("xl/media/") and n.endswith(".png") for n in names)
    book = load_workbook(io.BytesIO(data))
    assert "Summary" in book.sheetnames and "Plots BOL" in book.sheetnames


def test_pdf_has_text_and_images_and_is_deterministic(output) -> None:  # type: ignore[no-untyped-def]
    _, out = output
    data = render_pdf(out.document)
    assert render_pdf(out.document) == data
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Results by case" in text and "Violations" in text
    assert sum(len(page.images) for page in reader.pages) == 2  # BOL and EOL plots


def test_docx_opens_and_is_deterministic(output) -> None:  # type: ignore[no-untyped-def]
    _, out = output
    data = render_docx(out.document)
    assert render_docx(out.document) == data
    doc = Document(io.BytesIO(data))
    assert doc.core_properties.author == "Test User"
    assert doc.core_properties.created == WHEN
    assert len(doc.inline_shapes) == 2
    texts = [p.text for p in doc.paragraphs]
    assert any(t.startswith("Power budget over time") for t in texts)
    first = doc.tables[0]
    assert first.rows[0].cells[0].text == "Quantity"
    assert "BOL" in [c.text for c in first.rows[0].cells]


def test_docx_marks_missing_values(output) -> None:  # type: ignore[no-untyped-def]
    project = td_project(eta=None)
    result = time_domain_budget(project, synthetic_run(synthetic_env(1200.0, 20.0)))
    prov = make_provenance(project, scenario="s", user="u", generated_at=WHEN)
    doc = Document(io.BytesIO(render_docx(timeline_output(project, result, prov).document)))
    assert any("INCOMPLETE" in p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert "n/a" in cells


def test_series_csv_rows_match_the_steps(output) -> None:  # type: ignore[no-untyped-def]
    result, _ = output
    lines = series_csv(result, "bol").splitlines()
    assert lines[0].startswith("time_s,time_utc,mode,sunlight_ratio,load_w,demand_w,generation_w")
    assert len(lines) == result.steps + 1
    first = lines[1].split(",")
    assert first[0] == "0.000" and first[1] == "2026-01-01T00:00:00.000Z" and first[2] == "b"
    assert float(first[4]) == pytest.approx(100.0) and float(first[6]) == pytest.approx(100.0)
    assert float(first[11]) == pytest.approx(100.0)  # energy at the start: full battery
    assert len(series_csv(result, "bol", every=10).splitlines()) == result.steps // 10 + 2 - 1 + 0


def test_series_csv_leaves_missing_columns_empty() -> None:
    project = td_project(eta=None)
    result = time_domain_budget(project, synthetic_run(synthetic_env(600.0, 60.0)))
    row = series_csv(result, "bol").splitlines()[1].split(",")
    assert row[5] == "" and row[11] == ""  # demand and stored energy are n/a
    assert row[4] != ""  # the load is known


def test_json_summary(output) -> None:  # type: ignore[no-untyped-def]
    _, out = output
    data = json.loads(out.json_text)
    assert data["scenario"] == "synthetic" and data["load_basis"] == "nominal"
    assert set(data["cases"]) == {"bol", "eol"}
    assert data["cases"]["bol"]["capacity_wh"] == pytest.approx(100.0)
    assert data["provenance"]["user"] == "Test User"
    assert {v["code"] for v in data["violations"]} >= {"ORBIT_BALANCE_NEGATIVE"}
    assert data["limits"]["peak_power_w"] == 110.0


def test_write_outputs_names(output, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    _, out = output
    written = write_outputs([out], tmp_path, {"csv", "json", "docx"})
    assert sorted(p.name for p in written) == [
        "power_time_synthetic.docx",
        "power_time_synthetic.json",
        "power_time_synthetic_bol.csv",
        "power_time_synthetic_eol.csv",
        "power_time_synthetic_orbits.csv",
        "power_time_synthetic_violations.csv",
    ]


# ---- command line ---------------------------------------------------------------------------


@pytest.fixture
def eps(tmp_path: Path) -> Path:
    shutil.copytree(EXAMPLES / "cubesat_3u_eps", tmp_path / "p")
    return tmp_path / "p"


def test_cli_reports_findings_and_writes_every_file(
    eps: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out"
    code = main(["power-timeline", str(eps), "--out", str(out), *FIXED])
    text = capsys.readouterr().out
    assert code == 1  # the example's peak power is above its limit
    assert "PEAK_POWER_EXCEEDED" in text and "BOL:" in text and "EOL:" in text
    assert sorted(p.name for p in out.iterdir()) == [
        "power_time_one_day.docx",
        "power_time_one_day.json",
        "power_time_one_day.pdf",
        "power_time_one_day.xlsx",
        "power_time_one_day_bol.csv",
        "power_time_one_day_eol.csv",
        "power_time_one_day_orbits.csv",
        "power_time_one_day_violations.csv",
    ]


def test_cli_outputs_are_byte_identical(eps: Path, tmp_path: Path) -> None:
    for name in ("a", "b"):
        main(["power-timeline", str(eps), "--out", str(tmp_path / name), *FIXED])
    for path in sorted((tmp_path / "a").iterdir()):
        assert path.read_bytes() == (tmp_path / "b" / path.name).read_bytes(), path.name


def test_cli_single_case_and_series_thinning(eps: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    main(
        [
            "power-timeline",
            str(eps),
            "--case",
            "eol",
            "--report",
            "csv",
            "--series-every",
            "60",
            "--out",
            str(out),
            *FIXED,
        ]
    )
    assert sorted(p.name for p in out.iterdir()) == [
        "power_time_one_day_eol.csv",
        "power_time_one_day_orbits.csv",
        "power_time_one_day_violations.csv",
    ]
    assert len((out / "power_time_one_day_eol.csv").read_text().splitlines()) == 8640 // 60 + 1


def test_cli_margined_basis_and_phase_override(eps: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    main(
        [
            "power-timeline",
            str(eps),
            "--load-basis",
            "margined",
            "--phase",
            "launch",
            "--report",
            "json",
            "--out",
            str(out),
            *FIXED,
        ]
    )
    data = json.loads((out / "power_time_one_day.json").read_text())
    assert data["load_basis"] == "margined" and data["mission_phase"] == "launch"


def test_cli_without_a_power_system_says_what_is_missing(
    eps: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    (eps / "config" / "power_system.yaml").unlink()
    out = tmp_path / "out"
    code = main(["power-timeline", str(eps), "--report", "json", "--out", str(out), *FIXED])
    text = capsys.readouterr().out
    assert code == 0  # incomplete inputs are warnings, not findings
    assert "BOL: n/a" in text and "EOL: n/a" in text
    data = json.loads((out / "power_time_one_day.json").read_text())
    assert any(p["file"] == "config/power_system.yaml" for p in data["problems"])


def test_cli_strict_fails_on_warnings(eps: Path, tmp_path: Path) -> None:
    (eps / "config" / "power_system.yaml").unlink()
    code = main(
        [
            "power-timeline",
            str(eps),
            "--strict",
            "--report",
            "json",
            "--out",
            str(tmp_path / "o"),
            *FIXED,
        ]
    )
    assert code == 1


def test_cli_invalid_project_writes_nothing(
    eps: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    (eps / "project.yaml").unlink()
    out = tmp_path / "out"
    assert main(["power-timeline", str(eps), "--out", str(out)]) == 1
    assert not out.exists() and "nothing was written" in capsys.readouterr().out
