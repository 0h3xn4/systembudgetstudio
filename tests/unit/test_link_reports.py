"""Link outputs: report document, plots, XLSX/PDF/DOCX, CSV, JSON and the command line."""

from __future__ import annotations

import io
import json
import shutil
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from budget_cli.main import main
from budget_core.link.evaluate import link_pass_series, link_static_budget
from budget_core.plots.link import link_plot
from budget_core.plots.render import render_png
from budget_core.provenance import make_provenance
from budget_core.reports.docx import render_docx
from budget_core.reports.link_export import series_csv, static_csv
from budget_core.reports.pdf import render_pdf
from budget_core.reports.run import link_output, write_outputs
from budget_core.reports.xlsx import render_xlsx
from tests.link_helpers import link_project, pass_run

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
WHEN = datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)
FIXED = ["--user", "Test User", "--date", "2026-01-02T03:04:05Z"]


@pytest.fixture(scope="module")
def parts():  # type: ignore[no-untyped-def]
    project = link_project()
    run = pass_run(aos_s=100.0, los_s=400.0, duration_s=1000.0)
    static = link_static_budget(project)
    passes = link_pass_series(project, run)
    prov = make_provenance(project, scenario="synthetic", user="Test User", generated_at=WHEN)
    return project, static, passes, prov


def test_static_report_has_the_tables_and_no_banner(parts) -> None:  # type: ignore[no-untyped-def]
    project, static, _, prov = parts
    out = link_output(project, static, prov)
    assert out.prefix == "link_static"
    titles = [s.title for s in out.document.sections]
    assert titles == ["Summary", "Assumptions", "Problems", "Provenance"]
    table = out.document.sections[0].tables[0]
    assert table.rows[0][0] == "dl" and table.rows[0][2] == "slant"
    assert out.document.banner == ""
    assert [n for n, _ in out.csv_files] == ["link_static_table.csv"]


def test_pass_report_has_passes_volumes_and_a_plot(parts) -> None:  # type: ignore[no-untyped-def]
    project, static, passes, prov = parts
    out = link_output(project, static, prov, passes=passes)
    assert out.prefix == "link_passes_synthetic"
    titles = [s.title for s in out.document.sections]
    assert titles[:3] == ["Summary", "Passes", "Plot, dl"]
    volume = out.document.sections[1].tables[0]
    assert volume.rows[0][3] == pytest.approx(30.0)  # 100 kbit/s * 300 s = 30 Mbit
    pass_table = out.document.sections[1].tables[1]
    assert pass_table.rows[0][0] == "1" and pass_table.rows[0][3] == pytest.approx(300.0)
    figure = out.document.sections[2].figures[0]
    assert figure.png[:8] == b"\x89PNG\r\n\x1a\n"
    assert sorted(n for n, _ in out.csv_files) == [
        "link_passes_synthetic_dl.csv",
        "link_passes_synthetic_summary.csv",
        "link_passes_synthetic_table.csv",
    ]


def test_the_plot_has_one_series_per_pass_and_the_required_margin(parts) -> None:  # type: ignore[no-untyped-def]
    _, _, passes, _ = parts
    spec = link_plot(passes.series[0], 1000.0)
    assert [p.title for p in spec.panels] == ["Link margin", "Selected rate", "Elevation"]
    margin = spec.panels[0]
    assert len(margin.series) == 3 and [h.label for h in margin.hlines] == ["Required margin"]
    assert render_png(spec) == render_png(spec)


def test_renderers_are_deterministic_and_carry_the_figure(parts) -> None:  # type: ignore[no-untyped-def]
    project, static, passes, prov = parts
    doc = link_output(project, static, prov, passes=passes).document
    assert render_xlsx(doc) == render_xlsx(doc) and render_pdf(doc) == render_pdf(doc)
    assert render_docx(doc) == render_docx(doc)
    reader = PdfReader(io.BytesIO(render_pdf(doc)))
    assert "Data volume per link" in "\n".join(p.extract_text() for p in reader.pages)
    assert sum(len(p.images) for p in reader.pages) == 1
    assert len(Document(io.BytesIO(render_docx(doc))).inline_shapes) == 1


def test_static_csv_has_one_row_per_rate(parts) -> None:  # type: ignore[no-untyped-def]
    _, static, _, _ = parts
    lines = static_csv(static).splitlines()
    assert lines[0].startswith("link,direction,point,elevation_deg,range_m,frequency_hz,eirp_dbw")
    assert len(lines) == 1 + 3  # three data rates at the one static point
    assert lines[1].split(",")[2] == "slant" and lines[1].split(",")[-1] == "true"
    assert lines[3].split(",")[-1] == "false"  # 1 Mbit/s does not close


def test_series_csv_rows_are_the_samples_inside_the_pass(parts) -> None:  # type: ignore[no-untyped-def]
    _, _, passes, _ = parts
    lines = series_csv(passes, passes.series[0]).splitlines()
    assert len(lines) == 1 + 30
    first = lines[1].split(",")
    assert first[0] == "100.000" and first[1] == "2026-01-01T00:01:40.000Z" and first[2] == "1"
    assert lines[0].endswith(
        "margin_db_10000bps,margin_db_100000bps,margin_db_1000000bps,selected_rate_bps"
    )
    assert first[-1] == "100000.000000"


def test_json_summary(parts) -> None:  # type: ignore[no-untyped-def]
    project, static, passes, prov = parts
    data = json.loads(link_output(project, static, prov, passes=passes).json_text)  # type: ignore[arg-type]
    assert data["links"][0]["volume_bits"] == pytest.approx(3.0e7)
    assert data["links"][0]["passes"][0]["usable_s"] == pytest.approx(300.0)
    assert data["static"][0]["max_rate_bps"] == 100000.0
    assert data["provenance"]["user"] == "Test User"


def test_placeholders_give_na_cells_and_the_banner() -> None:
    project = link_project(ebn0=None)
    static = link_static_budget(project)
    prov = make_provenance(project, scenario="s", user="u", generated_at=WHEN)
    out = link_output(project, static, prov)
    assert "INCOMPLETE" in out.document.banner
    row = out.document.sections[0].tables[0].rows[0]
    assert row[10] is None  # C/N0
    assert static_csv(static).splitlines()[1].split(",")[11] == ""


def test_write_outputs_names(parts, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    project, static, _, prov = parts
    written = write_outputs([link_output(project, static, prov)], tmp_path, {"json", "csv", "docx"})
    assert sorted(p.name for p in written) == [
        "link_static.docx",
        "link_static.json",
        "link_static_table.csv",
    ]


# ---- command line ---------------------------------------------------------------------------


@pytest.fixture
def micro(tmp_path: Path) -> Path:
    shutil.copytree(EXAMPLES / "microsat_150kg", tmp_path / "p")
    return tmp_path / "p"


def test_run_link_budget_writes_the_static_files(micro: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    assert main(["run", str(micro), "--budget", "link", "--out", str(out), *FIXED]) == 0
    assert sorted(p.name for p in out.iterdir()) == [
        "link_static.docx",
        "link_static.json",
        "link_static.pdf",
        "link_static.xlsx",
        "link_static_table.csv",
    ]


def test_run_all_includes_links_only_when_the_project_has_them(micro: Path, tmp_path: Path) -> None:
    main(["run", str(micro), "--report", "json", "--out", str(tmp_path / "a"), *FIXED])
    assert "link_static.json" in {p.name for p in (tmp_path / "a").iterdir()}
    shutil.rmtree(micro / "links")
    main(["run", str(micro), "--report", "json", "--out", str(tmp_path / "b"), *FIXED])
    assert "link_static.json" not in {p.name for p in (tmp_path / "b").iterdir()}


def test_link_passes_writes_reports_and_prints_the_volume(
    micro: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out"
    code = main(
        [
            "link-passes",
            str(micro),
            "--scenario",
            "commissioning_day",
            "--out",
            str(out),
            "--report",
            "json",
            "--report",
            "csv",
            *FIXED,
        ]
    )
    text = capsys.readouterr().out
    assert code == 0 and "sband_down: 11 pass(es) over gs_north" in text and "MByte per day" in text
    names = {p.name for p in out.iterdir()}
    assert {
        "link_passes_commissioning_day.json",
        "link_passes_commissioning_day_xband_down.csv",
    } <= names


def test_link_passes_outputs_are_byte_identical(micro: Path, tmp_path: Path) -> None:
    for name in ("a", "b"):
        main(
            [
                "link-passes",
                str(micro),
                "--out",
                str(tmp_path / name),
                "--report",
                "xlsx",
                "--report",
                "json",
                *FIXED,
            ]
        )
    for path in sorted((tmp_path / "a").iterdir()):
        assert path.read_bytes() == (tmp_path / "b" / path.name).read_bytes(), path.name


def test_link_passes_without_links_or_with_a_wrong_scenario(
    micro: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert (
        main(["link-passes", str(micro), "--scenario", "nope", "--out", str(tmp_path / "o")]) == 1
    )
    assert "SCENARIO_UNKNOWN" in capsys.readouterr().out
    shutil.rmtree(micro / "links")
    assert main(["link-passes", str(micro), "--out", str(tmp_path / "o")]) == 1
    assert "no links" in capsys.readouterr().out


def test_placeholder_example_warns_and_strict_fails(tmp_path: Path) -> None:
    shutil.copytree(EXAMPLES / "cubesat_3u", tmp_path / "p")
    args = [
        "link-passes",
        str(tmp_path / "p"),
        "--out",
        str(tmp_path / "o"),
        "--report",
        "json",
        *FIXED,
    ]
    assert main(args) == 0
    assert main([*args, "--strict"]) == 1
    data = json.loads((tmp_path / "o" / "link_passes_one_day.json").read_text())
    assert data["links"][0]["volume_bits"] is None
