"""Mass budget report document, renderers and exports."""

from __future__ import annotations

import io
import json
import time
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from budget_core.io.project_loader import load_project
from budget_core.mass.static_mass import static_mass_budget
from budget_core.provenance import make_provenance
from budget_core.reports.document import NA, ReportDocument
from budget_core.reports.export import mass_csv, result_json
from budget_core.reports.mass_report import build_mass_report
from budget_core.reports.pdf import render_pdf
from budget_core.reports.xlsx import render_xlsx
from tests.helpers import write_valid_project

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
WHEN = datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)


def build(root: Path):  # type: ignore[no-untyped-def]
    result = load_project(root)
    assert result.project is not None
    mass = static_mass_budget(result.project)
    prov = make_provenance(result.project, user="Test User", generated_at=WHEN)
    return mass, prov, build_mass_report(result.project, mass, prov, result.problems)


def micro() -> ReportDocument:
    return build(EXAMPLES / "microsat_150kg")[2]


def test_sections_and_sheet_names() -> None:
    doc = micro()
    names = [s.sheet_name for s in doc.sections]
    assert names == [
        "Summary",
        "Phase launch",
        "Phase bol",
        "Phase eol",
        "Assumptions",
        "Problems",
        "Provenance",
    ]
    assert all(len(n) <= 31 for n in names)
    assert doc.title.startswith("Mass budget")


def test_placeholders_give_na_and_a_banner_but_geometry_numbers_are_real() -> None:
    doc = micro()
    assert doc.banner.startswith("INCOMPLETE")
    summary = doc.sections[0].tables[0]
    row = summary.rows[0]  # launch
    assert isinstance(row[1], float)  # nominal total is computed
    assert row[2] is None and row[3] is None  # margined totals need placeholders
    assert all(isinstance(v, float) for v in row[4:7])  # centre of gravity x, y, z
    limits = doc.sections[0].tables[1]
    assert limits.rows[0][-1] == NA or limits.rows[0][-1] == "n/a"


def test_phase_section_has_items_subsystems_totals_and_inertia() -> None:
    doc = micro()
    phase = next(s for s in doc.sections if s.sheet_name == "Phase launch")
    titles = [t.title for t in phase.tables]
    assert titles == ["Items", "By subsystem", "Totals and centre of gravity", "Inertia tensor"]
    items = phase.tables[0]
    assert {r[2] for r in items.rows} == {"unit", "expendable"}
    assert any(r[0] == "adapter" for r in items.rows)  # jettisoned later
    eol = next(s for s in doc.sections if s.sheet_name == "Phase eol")
    assert not any(r[0] == "adapter" for r in eol.tables[0].rows)


def test_fully_sourced_project_has_no_banner_and_checks_the_limit(tmp_path: Path) -> None:
    write_valid_project(tmp_path)
    _, _, doc = build(tmp_path)
    assert doc.banner == ""
    limits = doc.sections[0].tables[1]
    assert limits.rows[0][-1] == "ok"


def test_xlsx_and_pdf_render_and_are_clock_independent() -> None:
    doc = micro()
    first_x, first_p = render_xlsx(doc), render_pdf(doc)
    time.sleep(1.2)
    assert render_xlsx(doc) == first_x and render_pdf(doc) == first_p
    wb = load_workbook(io.BytesIO(first_x))
    text = "\n".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "MASS-COG" in text and "MASS-PARALLEL" in text and "Test User" in text
    assert first_p.startswith(b"%PDF-")


def test_json_export_is_deterministic_and_complete() -> None:
    mass, prov, _ = build(EXAMPLES / "microsat_150kg")
    text = result_json(mass, prov)
    assert text == result_json(mass, prov)
    data = json.loads(text)
    assert [p["phase"] for p in data["phases"]] == ["launch", "bol", "eol"]
    launch = data["phases"][0]
    assert launch["cog"]["position_m"][2] > 0 and launch["system_margined_kg"] is None


def test_csv_export_lists_items_with_positions() -> None:
    mass, _, _ = build(EXAMPLES / "microsat_150kg")
    text = mass_csv(mass.phases[0])
    lines = text.strip().split("\n")
    assert lines[0] == (
        "item_id,name,kind,subsystem,maturity,mass_kg,margin_ratio,margined_mass_kg,x_m,y_m,z_m"
    )
    assert len(lines) == 1 + len(mass.phases[0].rows) and "\r" not in text
    assert any(line.startswith("propellant,Propellant,expendable") for line in lines)
