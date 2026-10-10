"""Content dumps of report files for golden tests (compressed bytes differ between zlib builds)."""

from __future__ import annotations

import io
import os
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

GOLDEN_DIR = Path(__file__).parent / "golden"


# A decimal number that is not part of a version string or a date ("0.1.0", "1.2.3").
_FLOAT = re.compile(r"(?<![\w.])-?\d+(?:\.\d+(?:[eE][+-]?\d+)?|[eE][+-]?\d+)(?![\w.])")


def stable_floats(text: str, digits: int = 9, zero_below: float = 1e-6) -> str:
    """Text with every decimal number rounded to `digits` significant digits and every number
    smaller than `zero_below` written as 0.0. Results of linear algebra (a different BLAS build
    changes the last digits, and the residual of a solved system is pure rounding noise) are
    compared at the precision the physics supports, not bit for bit."""

    def fix(match: re.Match[str]) -> str:
        value = float(match.group(0))
        return "0.0" if abs(value) < zero_below else repr(float(f"{value:.{digits}g}"))

    return _FLOAT.sub(fix, text)


def dump_xlsx(
    data: bytes,
    max_rows: int | None = None,
    significant_digits: int | None = None,
    zero_below: float | None = None,
) -> str:
    wb = load_workbook(io.BytesIO(data))
    lines: list[str] = []
    for ws in wb:
        lines.append(f"## sheet {ws.title} rows={ws.max_row} cols={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if max_rows is not None and i > max_rows:
                lines.append("...")
                break
            if significant_digits is not None:  # sums over many steps differ in the last digits
                row = tuple(
                    float(f"{v:.{significant_digits}g}") if isinstance(v, float) else v for v in row
                )
            if zero_below is not None:
                row = tuple(0.0 if isinstance(v, float) and abs(v) < zero_below else v for v in row)
            cells = ["" if v is None else repr(v) for v in row]
            if any(cells):
                lines.append(f"{i}: " + " | ".join(cells))
    return "\n".join(lines) + "\n"


def dump_pdf(data: bytes, max_pages: int | None = None) -> str:
    reader = PdfReader(io.BytesIO(data))
    lines = [f"pages={len(reader.pages)}"]
    for number, page in enumerate(reader.pages, 1):
        if max_pages is not None and number > max_pages:
            lines.append("...")
            break
        lines.append(f"## page {number}")
        lines.append((page.extract_text() or "").strip())
    return "\n".join(lines) + "\n"


def dump_docx(data: bytes) -> str:
    """Paragraph texts and table cells in document order, plus the picture count."""
    doc = Document(io.BytesIO(data))
    lines = [f"pictures={len(doc.inline_shapes)}"]
    body = doc.element.body
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    for child in body.iterchildren():
        if child in paragraphs and paragraphs[child].text.strip():
            lines.append(paragraphs[child].text)
        elif child in tables:
            lines.append("## table")
            for row in tables[child].rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines) + "\n"


def check_golden(name: str, actual: str) -> None:
    path = GOLDEN_DIR / name
    if os.environ.get("UPDATE_GOLDEN") == "1":
        path.write_text(actual, encoding="utf-8", newline="\n")
        return
    assert path.exists(), f"missing golden file {name}; run with UPDATE_GOLDEN=1"
    assert actual == path.read_text(encoding="utf-8"), (
        f"{name} differs; if intended, regenerate with UPDATE_GOLDEN=1 and review the diff"
    )
